"""Разбор ДПО-программы из DOCX в :class:`DpoProgram`.

Идёт по блокам документа в порядке следования (абзацы + таблицы), нарезает на разделы по стилям
заголовков (``Heading N``) и складывает тело/таблицы в текущий раздел. Никакой интерпретации —
только структурный перенос, чтобы дальше слот-аксессоры отдавали текст «как есть».
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import docx
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from .models import DpoProgram, ProgramSection, ProgramTable

#: «1.11. Заголовок» / «2. ЗАГОЛОВОК» → (номер, заголовок).
_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(.*\S)\s*$")


def parse_dpo_program(path: str | Path) -> DpoProgram:
    """Разбирает DOCX действующей ДПО-программы в типизированную модель."""
    document = docx.Document(str(path))
    sections: list[ProgramSection] = []
    current: ProgramSection | None = None
    title = ""

    for block in _iter_blocks(document):
        if isinstance(block, DocxTable):
            if current is not None:
                current.tables.append(_table_model(block))
            continue

        text = block.text.strip()
        if not text:
            continue

        if _is_heading(block):
            number, heading_title = _split_heading(text)
            current = ProgramSection(number=number, title=heading_title, level=_heading_level(block))
            sections.append(current)
        elif current is None:
            # Титульная преамбула до первого заголовка — берём первую строку как имя документа.
            if not title:
                title = text
        else:
            current.paragraphs.append(text)

    return DpoProgram(title=title, sections=sections)


def _iter_blocks(document: Any) -> Iterator[Paragraph | DocxTable]:
    """Абзацы и таблицы тела документа в порядке следования (python-docx их не смешивает сам)."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def _is_heading(paragraph: Paragraph) -> bool:
    style = _style_name(paragraph)
    return style.startswith("heading") or style.startswith("title") or style.startswith("заголов")


def _heading_level(paragraph: Paragraph) -> int:
    match = re.search(r"(\d+)", _style_name(paragraph))
    return int(match.group(1)) if match else 0


def _style_name(paragraph: Paragraph) -> str:
    style = paragraph.style
    return (style.name if style and style.name else "").strip().lower()


def _split_heading(text: str) -> tuple[str, str]:
    match = _HEADING_RE.match(text)
    if match:
        return match.group(1), match.group(2).strip()
    return "", text


def _table_model(table: DocxTable) -> ProgramTable:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ProgramTable()
    return ProgramTable(header=rows[0], rows=rows[1:])
