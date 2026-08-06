"""Выборочная сборка РПД на реальном ДПО-шаблоне (повышение квалификации).

В отличие от полной очистки тела: идём по разделам шаблона и заменяем ТОЛЬКО те, что должны прийти
из УП / генерации / пробела; разделы-boilerplate (нормативка, документ, условия инвалидов) оставляем
как в шаблоне. Категорию раздела определяет :func:`dpo_spec.category_for`.
"""

from __future__ import annotations

import re
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .build import RpdBuild, RpdSection, build_rpd_sections
from .ingest import _is_heading, _split_heading
from .template_profiles import get_template_profile
from .up_models import UpPlan

_Content = tuple[list[str], list[tuple[list[str], list[list[str]]]]]


def render_rpd_pk(up: UpPlan, template_path: str | Path | None = None, llm_client: Any = None) -> bytes:
    """Удобная обёртка: собрать черновик РПД из УП (+ проза LLM) и сразу отрендерить в DOCX."""
    build = build_rpd_sections(up, template_path, llm_client)
    return render_build(build, template_path)


def render_build(build: RpdBuild, template_path: str | Path | None = None) -> bytes:
    """Кладёт (возможно отредактированный) черновик РПД в DOCX по шаблону: editable-разделы заменяет,
    boilerplate (форму ДПО) держит как в шаблоне. Комментарии вырезаются."""
    profile = get_template_profile(build.template_profile_id)
    path = profile.template_path(template_path)
    document = docx.Document(str(path))
    body = document.element.body
    sectpr = qn("w:sectPr")
    para_tag = qn("w:p")
    table_tag = qn("w:tbl")

    elements = [child for child in body.iterchildren() if child.tag != sectpr]
    spans = _section_spans(document, elements, para_tag)

    for heading, number, title, body_els in spans:
        section = build.section(number, title)
        if section is None:
            continue  # boilerplate / заголовок-контейнер — держим шаблон
        table_prototypes = [deepcopy(element) for element in body_els if element.tag == table_tag]
        for element in body_els:
            # A section break is stored inside a paragraph's pPr. Removing that apparently empty
            # paragraph silently collapses portrait/landscape sections in the resulting DOCX.
            if not _has_section_break(element):
                body.remove(element)
        _insert_after(document, heading, _content_of(section), table_prototypes)

    # Обложка и паспорт находятся вне секций с заголовками, поэтому заполняются отдельным контрактом.
    profile.bind_document(document, build)

    buffer = BytesIO()
    document.save(buffer)
    return _drop_comments(buffer.getvalue())


def _content_of(section: RpdSection) -> _Content:
    if section.kind == "table" and section.table_rows:
        return ([], [(section.table_header, section.table_rows)])
    return (list(section.paragraphs), [])


# Части-комментарии и их разметка, которые нельзя оставлять в готовом РПД (в шаблоне это заметки методиста).
_COMMENT_PARTS = {
    "word/comments.xml",
    "word/commentsExtended.xml",
    "word/commentsIds.xml",
    "word/commentsExtensible.xml",
    "word/people.xml",
}
_COMMENT_ELEMENT_RE = re.compile(r"<w:comment(?:RangeStart|RangeEnd|Reference)\b[^>]*/>")
_COMMENT_OVERRIDE_RE = re.compile(r"<Override\b[^>]*comments[^>]*/>")
_COMMENT_REL_RE = re.compile(r"<Relationship\b[^>]*comments[^>]*/>")


def _drop_comments(data: bytes) -> bytes:
    """Убирает все комментарии из DOCX: части comments*.xml, их привязки в тексте, rels и content-types."""
    source = zipfile.ZipFile(BytesIO(data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename in _COMMENT_PARTS:
                continue
            payload = source.read(item.filename)
            if item.filename == "word/document.xml":
                payload = _COMMENT_ELEMENT_RE.sub("", payload.decode("utf-8")).encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                payload = _COMMENT_OVERRIDE_RE.sub("", payload.decode("utf-8")).encode("utf-8")
            elif item.filename.endswith("document.xml.rels"):
                payload = _COMMENT_REL_RE.sub("", payload.decode("utf-8")).encode("utf-8")
            target.writestr(item, payload)
    return out.getvalue()


def _section_spans(document: Any, elements: list[Any], para_tag: str) -> list[tuple[Any, str, str, list[Any]]]:
    spans: list[tuple[Any, str, str, list[Any]]] = []
    heading: Any = None
    number = title = ""
    body_els: list[Any] = []

    def flush() -> None:
        if heading is not None:
            spans.append((heading, number, title, body_els))

    for element in elements:
        if element.tag == para_tag and _is_heading(Paragraph(element, document)):
            flush()
            heading = element
            number, title = _split_heading(Paragraph(element, document).text.strip())
            body_els = []
        elif heading is not None:
            body_els.append(element)
    flush()
    return spans


def _insert_after(document: Any, anchor: Any, content: _Content, table_prototypes: list[Any] | None = None) -> None:
    paragraphs, tables = content
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        anchor.addnext(paragraph._p)
        anchor = paragraph._p
    for header, rows in tables:
        prototype = _matching_table_prototype(document, table_prototypes or [], max(len(header), 1))
        grid = _build_table(document, header, rows, prototype)
        anchor.addnext(grid._tbl)
        anchor = grid._tbl


def _build_table(document: Any, header: list[str], rows: list[list[str]], prototype: Any = None) -> Any:
    ncols = max(len(header), max((len(r) for r in rows), default=0), 1)
    if prototype is not None:
        # Clone the template table so widths, borders, shading, alignment and fonts survive.
        table_element = deepcopy(prototype)
        existing_rows = list(table_element.tr_lst)
        if existing_rows:
            data_prototype = deepcopy(existing_rows[1] if len(existing_rows) > 1 else existing_rows[0])
            for row_element in existing_rows[1:]:
                table_element.remove(row_element)
            for _row in rows:
                table_element.append(deepcopy(data_prototype))
            grid = Table(table_element, document)
            _populate_table(grid, header, rows, ncols)
            return grid

    grid = document.add_table(rows=len(rows) + 1, cols=ncols)
    try:
        grid.style = "Table Grid"
    except (KeyError, ValueError):
        pass  # стиля может не быть в каркасе — границы всё равно ставим явно
    _set_grid_borders(grid)
    _populate_table(grid, header, rows, ncols)
    return grid


def _matching_table_prototype(document: Any, prototypes: list[Any], ncols: int) -> Any | None:
    """Pick the table in a section whose visible schema matches generated content."""
    for prototype in prototypes:
        table = Table(prototype, document)
        if table.rows and len(table.rows[0].cells) == ncols:
            return prototype
    return None


def _populate_table(grid: Any, header: list[str], rows: list[list[str]], ncols: int) -> None:
    for col in range(ncols):
        _replace_cell_text(grid.rows[0].cells[col], header[col] if col < len(header) else "")
    for row_index, row in enumerate(rows, start=1):
        for col in range(ncols):
            _replace_cell_text(grid.rows[row_index].cells[col], row[col] if col < len(row) else "")


def _replace_cell_text(cell: Any, text: str) -> None:
    """Replace text while retaining the template paragraph and first-run formatting."""
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _has_section_break(element: Any) -> bool:
    return any(node.tag == qn("w:sectPr") for node in element.iter())


def _set_grid_borders(grid: Any) -> None:
    """Явные одинарные границы: шаблонного стиля «Table Grid» может не быть → иначе таблица без рамок."""
    tbl_pr = grid._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)
