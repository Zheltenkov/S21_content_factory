"""Explicit template profiles for deterministic RPD rendering.

The profile is the seam between the domain build and a concrete DOCX family.  It owns template
recognition, table schemas and cover/passport bindings.  Adding a second document family therefore
does not add conditionals to the orchestration or renderer.
"""

from __future__ import annotations

import hashlib
import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import TYPE_CHECKING, Any

import docx
from docx.enum.section import WD_ORIENT
from pydantic import BaseModel, Field

from .dpo_spec import Cat, category_for
from .template import default_pk_skeleton_path, load_rpd_template
from .up_models import UpPlan

if TYPE_CHECKING:
    from .build import RpdBuild


TableContent = tuple[list[str], list[list[str]]]
BUILTIN_PK_PROFILE_ID = "dpo_pk_v1"


class UnsupportedTemplateProfileError(ValueError):
    """Raised when a persisted build references a profile unknown to this deployment."""


class TemplateCapabilityReport(BaseModel):
    """Machine-readable evidence that a DOCX satisfies a profile contract."""

    profile_id: str
    program_type: str
    template_sha256: str
    supported: bool
    section_count: int = 0
    missing_bindings: list[str] = Field(default_factory=list)
    missing_table_labels: list[str] = Field(default_factory=list)
    missing_section_keys: list[str] = Field(default_factory=list)
    missing_table_schemas: list[str] = Field(default_factory=list)
    duplicate_section_numbers: list[str] = Field(default_factory=list)
    section_orientations: list[str] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)


class TemplateProfile(ABC):
    """Small interface hiding all format-specific DOCX and table knowledge."""

    profile_id: str
    program_type: str

    @abstractmethod
    def template_path(self, override: str | Path | None = None) -> Path:
        """Resolve the concrete DOCX used for this build."""

    @abstractmethod
    def inspect(self, path: Path) -> TemplateCapabilityReport:
        """Validate whether a template exposes the required bindings and table labels."""

    def category_for(self, number: str, title: str) -> Cat:
        """Return the deterministic fallback category for a section."""
        return category_for(number, title)

    @abstractmethod
    def table_for(self, category: Cat, up: UpPlan) -> TableContent | None:
        """Map normalized curriculum data to a profile-specific table schema."""

    @abstractmethod
    def bind_document(self, document: Any, build: RpdBuild) -> None:
        """Apply profile-specific cover and passport bindings to an open DOCX document."""


class BuiltinPkProfile(TemplateProfile):
    """Production profile for the bundled professional-development (ПК) template."""

    profile_id = BUILTIN_PK_PROFILE_ID
    program_type = "повышение квалификации"

    _required_bindings = {
        "program_title": "[НАИМЕНОВАНИЕ ПРОГРАММЫ]",
        "organization_name": "[ПОЛНОЕ НАИМЕНОВАНИЕ ОБРАЗОВАТЕЛЬНОЙ ОРГАНИЗАЦИИ]",
        "review_body": "[наименование коллегиального органа]",
        "leader_position": "[должность руководителя]",
        "leader_name": "[Ф.И.О.]",
        "city": "[ГОРОД]",
        "year": "[ГОД]",
    }
    _required_table_labels = {
        "Категория слушателей",
        "Объем программы",
        "Форма обучения",
        "Итоговая аттестация",
        "Выдаваемый документ",
    }
    _required_section_keys = {
        "2.2 Профессиональные компетенции и их качественное изменение",
        "3.1 Учебный план",
        "3.2 Календарный учебный график",
        "3.3.1 Пример оформления учебного модуля",
        "4.3 Паспорт оценочных материалов",
        "5.5 Условия для инвалидов и лиц с ограниченными возможностями здоровья",
    }
    _required_table_schemas = {
        "компетенции": (
            "Код",
            "Профессиональная компетенция",
            "Исходный уровень",
            "Качественное изменение",
            "После обучения слушатель способен",
            "Основание",
            "Оценка",
        ),
        "учебный план": (
            "№",
            "Наименование модуля",
            "Всего",
            "Лекции",
            "Практические занятия",
            "Самостоятельная работа",
            "Контроль",
        ),
        "календарный график": ("Компонент", "Период", "Часы", "Форма организации", "Контроль"),
        "рабочая программа модуля": ("№", "Тема и краткое содержание", "Вид работы", "Часы", "Методы", "Контроль"),
    }

    def template_path(self, override: str | Path | None = None) -> Path:
        return Path(override) if override is not None else default_pk_skeleton_path()

    def inspect(self, path: Path) -> TemplateCapabilityReport:
        data = path.read_bytes()
        document = docx.Document(str(path))
        text = "\n".join(paragraph.text for paragraph in _all_paragraphs(document))
        missing_bindings = [name for name, marker in self._required_bindings.items() if marker not in text]
        table_labels = {
            row.cells[0].text.strip()
            for table in document.tables
            for row in table.rows
            if row.cells
        }
        missing_table_labels = sorted(self._required_table_labels - table_labels)
        template = load_rpd_template(path)
        section_keys = {f"{section.number} {section.title}".strip() for section in template.sections}
        missing_section_keys = sorted(self._required_section_keys - section_keys)
        table_schemas = {
            tuple(cell.text.strip() for cell in table.rows[0].cells)
            for table in document.tables
            if table.rows
        }
        missing_table_schemas = [
            name for name, schema in self._required_table_schemas.items() if schema not in table_schemas
        ]
        number_counts: dict[str, int] = {}
        for section in template.sections:
            if section.number:
                number_counts[section.number] = number_counts.get(section.number, 0) + 1
        duplicate_section_numbers = sorted(number for number, count in number_counts.items() if count > 1)
        section_orientations = [
            "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait"
            for section in document.sections
        ]
        warnings: list[str] = []
        if not any(self.category_for(section.number, section.title) is Cat.UP_PLAN for section in template.sections):
            warnings.append("Не найден раздел учебного плана.")
        if len(document.sections) != 4 or "landscape" not in section_orientations:
            warnings.append("Нарушен контракт секций DOCX: ожидаются 4 секции, включая альбомную.")
        supported = not any(
            (
                missing_bindings,
                missing_table_labels,
                missing_section_keys,
                missing_table_schemas,
                duplicate_section_numbers,
                warnings,
            )
        )
        return TemplateCapabilityReport(
            profile_id=self.profile_id,
            program_type=self.program_type,
            template_sha256=hashlib.sha256(data).hexdigest(),
            supported=supported,
            section_count=len(template.sections),
            missing_bindings=missing_bindings,
            missing_table_labels=missing_table_labels,
            missing_section_keys=missing_section_keys,
            missing_table_schemas=missing_table_schemas,
            duplicate_section_numbers=duplicate_section_numbers,
            section_orientations=section_orientations,
            warnings=warnings,
        )

    def table_for(self, category: Cat, up: UpPlan) -> TableContent | None:
        handlers = {
            Cat.UP_PLAN: self._curriculum_table,
            Cat.UP_GRAFIK: self._schedule_table,
            Cat.UP_MODULES: self._module_table,
            Cat.UP_COMP: self._competency_table,
        }
        handler = handlers.get(category)
        return handler(up) if handler is not None else None

    def bind_document(self, document: Any, build: RpdBuild) -> None:
        fields = build.document_fields
        replacements = {
            self._required_bindings["program_title"]: build.program_title,
            self._required_bindings["organization_name"]: fields.organization_name,
            self._required_bindings["review_body"]: fields.review_body,
            self._required_bindings["leader_position"]: fields.leader_position,
            self._required_bindings["leader_name"]: fields.leader_name,
            self._required_bindings["city"]: fields.city,
            self._required_bindings["year"]: fields.year,
        }
        active_replacements = {key: value for key, value in replacements.items() if value.strip()}
        for paragraph in _all_paragraphs(document):
            text = paragraph.text
            updated = text
            for source, target in active_replacements.items():
                updated = updated.replace(source, target)
            if fields.protocol_number and fields.protocol_date:
                updated = re.sub(
                    r"Протокол\s+№\s+___\s+от\s+«___»\s+________\s+20__\s+г\.",
                    f"Протокол № {fields.protocol_number} от {fields.protocol_date}",
                    updated,
                )
            if fields.order_number and fields.order_date:
                updated = re.sub(
                    r"Приказ\s+№\s+___\s+от\s+«___»\s+________\s+20__\s+г\.",
                    f"Приказ № {fields.order_number} от {fields.order_date}",
                    updated,
                )
            updated = updated.replace("[дд.мм.гггг]", "")
            if updated.strip() == "Пример: «Цифровые инструменты в профессиональной деятельности»":
                updated = ""
            if updated != text:
                _replace_paragraph_text(paragraph, updated)
        self._fill_program_parameters(document, build.total_hours, fields)
        self._clear_change_log_example(document)

    @staticmethod
    def _curriculum_table(up: UpPlan) -> TableContent | None:
        if not up.modules:
            return None
        header = [
            "№",
            "Наименование модуля",
            "Всего",
            "Лекции",
            "Практические занятия",
            "Самостоятельная работа",
            "Контроль",
        ]
        rows = [[str(index), module.title, module.hours, "", module.hours, "", module.control()] for index, module in enumerate(up.modules, start=1)]
        return header, rows

    @staticmethod
    def _schedule_table(up: UpPlan) -> TableContent | None:
        if not up.modules:
            return None
        header = ["Компонент", "Период", "Часы", "Форма организации", "Контроль"]
        return header, [[module.title, "", module.hours, module.fmt(), module.control()] for module in up.modules]

    @staticmethod
    def _module_table(up: UpPlan) -> TableContent | None:
        if not up.modules:
            return None
        header = ["№", "Тема и краткое содержание", "Вид работы", "Часы", "Методы", "Контроль"]
        rows: list[list[str]] = []
        index = 1
        for module in up.modules:
            for row in module.rows:
                theme = f"[{module.title}] {row.name}".strip()
                if row.description:
                    theme = f"{theme}: {row.description}"
                rows.append([str(index), theme, row.fmt, row.hours, "", row.control()])
                index += 1
            # Комментарии методиста к шаблону требуют подтверждение из LMS для каждого модуля.
            # Маркер живёт в редактируемой таблице и блокирует финальный экспорт до замены.
            rows.append(
                [
                    "",
                    f"[Методист: вставьте скрин заданий и контролей модуля «{module.title}» с платформы/LMS]",
                    "",
                    "",
                    "",
                    "",
                ]
            )
        return (header, rows) if rows else None

    @staticmethod
    def _competency_table(up: UpPlan) -> TableContent | None:
        competency_rows = up.competency_rows()
        if not competency_rows:
            return None
        header = [
            "Код",
            "Профессиональная компетенция",
            "Исходный уровень",
            "Качественное изменение",
            "После обучения слушатель способен",
            "Основание",
            "Оценка",
        ]
        return header, [[code, name, "", "", after, "", ""] for code, name, after in competency_rows]

    @staticmethod
    def _fill_program_parameters(document: Any, total_hours: str, fields: Any) -> None:
        values = {
            "Категория слушателей": fields.audience,
            "Объем программы": f"{total_hours} академических часов" if total_hours else "",
            "Форма обучения": fields.delivery_mode,
            "Итоговая аттестация": fields.final_attestation,
            "Выдаваемый документ": fields.issued_document,
        }
        for table in document.tables:
            if table.rows and len(table.rows[0].cells) >= 2 and table.rows[0].cells[0].text.strip() == "Параметр":
                _replace_paragraph_text(table.rows[0].cells[1].paragraphs[0], "Значение")
            for row in table.rows:
                if len(row.cells) < 2:
                    continue
                value = values.get(row.cells[0].text.strip(), "")
                if value:
                    _replace_paragraph_text(row.cells[1].paragraphs[0], value)

    @staticmethod
    def _clear_change_log_example(document: Any) -> None:
        expected = {"Версия", "Дата", "Раздел", "Содержание изменения", "Основание", "Ответственный"}
        for table in document.tables:
            if not table.rows:
                continue
            header = {cell.text.strip() for cell in table.rows[0].cells}
            if expected.issubset(header):
                for row in table.rows[1:]:
                    for cell in row.cells:
                        _replace_paragraph_text(cell.paragraphs[0], "")


_PROFILES: dict[str, TemplateProfile] = {BUILTIN_PK_PROFILE_ID: BuiltinPkProfile()}


def get_template_profile(profile_id: str = BUILTIN_PK_PROFILE_ID) -> TemplateProfile:
    """Resolve a registered profile or fail explicitly instead of guessing a document family."""
    try:
        return _PROFILES[profile_id]
    except KeyError as exc:
        raise UnsupportedTemplateProfileError(f"Unsupported RPD template profile: {profile_id}") from exc


def registered_template_profiles() -> tuple[str, ...]:
    """Return stable profile identifiers available in this deployment."""
    return tuple(_PROFILES)


def detect_template_profile(path: str | Path) -> tuple[TemplateProfile, TemplateCapabilityReport]:
    """Recognize a DOCX by capability contracts; never guess when no profile matches."""
    resolved = Path(path)
    reports = [(profile, profile.inspect(resolved)) for profile in _PROFILES.values()]
    matches = [(profile, report) for profile, report in reports if report.supported]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        identifiers = ", ".join(profile.profile_id for profile, _report in matches)
        raise UnsupportedTemplateProfileError(f"Ambiguous RPD template profiles: {identifiers}")
    evidence = "; ".join(
        f"{report.profile_id}: bindings={report.missing_bindings}, tables={report.missing_table_labels}"
        for _profile, report in reports
    )
    raise UnsupportedTemplateProfileError(f"Unsupported RPD template: {evidence}")


def _all_paragraphs(document: Any) -> list[Any]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_table_paragraphs(table))
    for section in document.sections:
        for container in (section.header, section.footer):
            paragraphs.extend(container.paragraphs)
            for table in container.tables:
                paragraphs.extend(_table_paragraphs(table))
    return paragraphs


def _table_paragraphs(table: Any) -> list[Any]:
    paragraphs: list[Any] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for nested in cell.tables:
                paragraphs.extend(_table_paragraphs(nested))
    return paragraphs


def _replace_paragraph_text(paragraph: Any, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


__all__ = [
    "BUILTIN_PK_PROFILE_ID",
    "BuiltinPkProfile",
    "TemplateCapabilityReport",
    "TemplateProfile",
    "UnsupportedTemplateProfileError",
    "detect_template_profile",
    "get_template_profile",
    "registered_template_profiles",
]
