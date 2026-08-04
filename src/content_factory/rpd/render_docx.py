"""Сборка выходного РПД в DOCX из :class:`RpdDraft`.

Временный подход (каркас = образец `_итог.docx`): открываем каркас — чтобы унаследовать его стили
(заголовки/шрифты Школы 21) и определения таблиц, — затем ПОЛНОСТЬЮ очищаем тело и пересобираем его
разделами из черновика. Полная очистка тела гарантирует, что контент образца не протечёт в чужую РПД.
При приходе настоящего ДПО-шаблона достаточно подменить файл-каркас — код рендера не меняется.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import docx
from docx.oxml.ns import qn

from .models import ProgramTable
from .transfer import FilledSection, RpdDraft

_RPD_TITLE = "РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ"


def render_rpd_docx(skeleton_path: str | Path, draft: RpdDraft, program_title: str = "") -> bytes:
    """Собирает РПД: стили берём из каркаса, тело пересобираем из ``draft``. Возвращает байты DOCX."""
    document = docx.Document(str(skeleton_path))
    _clear_body(document)

    document.add_heading(_RPD_TITLE, level=0)
    if program_title.strip():
        document.add_paragraph(program_title.strip())

    for section in draft.sections:
        _add_section(document, section)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _clear_body(document: Any) -> None:
    """Удаляет все абзацы и таблицы тела, сохраняя ``w:sectPr`` (параметры страницы) и стили."""
    body = document.element.body
    sectpr = qn("w:sectPr")
    for child in list(body.iterchildren()):
        if child.tag != sectpr:
            body.remove(child)


def _add_section(document: Any, section: FilledSection) -> None:
    heading = _heading_text(section)
    if heading:
        document.add_heading(heading, level=_clamp_level(section.level))
    if section.is_container:
        return
    for paragraph in section.paragraphs:
        document.add_paragraph(paragraph)
    for table in section.tables:
        _add_table(document, table)


def _heading_text(section: FilledSection) -> str:
    if section.number and section.title:
        return f"{section.number}. {section.title}"
    return section.number or section.title


def _clamp_level(level: int) -> int:
    if level < 1:
        return 1
    return min(level, 9)


def _add_table(document: Any, table: ProgramTable) -> None:
    rows = ([table.header] if table.header else []) + table.rows
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    if ncols < 1:
        return
    grid = document.add_table(rows=len(rows), cols=ncols)
    try:
        grid.style = "Table Grid"
    except (KeyError, ValueError):
        pass  # стиль отсутствует в каркасе — оставляем таблицу без явного стиля
    for row_index, row in enumerate(rows):
        for col_index in range(ncols):
            grid.rows[row_index].cells[col_index].text = row[col_index] if col_index < len(row) else ""
