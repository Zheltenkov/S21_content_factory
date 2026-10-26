"""XLSX: перевод строковых ячеек правкой XML внутри книги.

Книга не пересобирается openpyxl: пересборка теряет диаграммы, картинки и часть оформления.
Правятся только общие строки и встроенный текст листов, остальные части архива копируются как есть.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from content_factory.translation.errors import DocumentFormatError
from content_factory.translation.handlers.xlsx import XlsxHandler


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Project name"
    sheet["B1"] = "Hours"
    sheet["A2"] = "Intro to Machine Learning"
    sheet["B2"] = 36
    sheet["A3"] = "Decision boundaries"
    sheet["B3"] = 12
    sheet["B4"] = "=SUM(B2:B3)"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cells(content: bytes) -> list[object]:
    sheet = load_workbook(BytesIO(content)).active
    return [cell.value for row in sheet.iter_rows() for cell in row]


@pytest.fixture
def handler() -> XlsxHandler:
    return XlsxHandler()


def test_text_cells_become_units(handler: XlsxHandler) -> None:
    texts = [u.text for u in handler.extract(_xlsx_bytes())]

    assert "Project name" in texts
    assert "Intro to Machine Learning" in texts
    assert "Decision boundaries" in texts


def test_numbers_and_formulas_are_not_units(handler: XlsxHandler) -> None:
    texts = [u.text for u in handler.extract(_xlsx_bytes())]

    assert "36" not in texts and "12" not in texts
    assert not any("SUM" in text for text in texts), "формула не текст для перевода"


def test_without_translations_every_cell_survives(handler: XlsxHandler) -> None:
    source = _xlsx_bytes()

    assert _cells(handler.apply(source, {})) == _cells(source)


def test_a_translated_cell_replaces_only_itself(handler: XlsxHandler) -> None:
    source = _xlsx_bytes()
    unit = next(u for u in handler.extract(source) if u.text == "Intro to Machine Learning")

    values = _cells(handler.apply(source, {unit.unit_id: "Введение в машинное обучение"}))

    assert "Введение в машинное обучение" in values
    assert "Decision boundaries" in values, "соседняя ячейка не тронута"
    assert 36 in values, "число осталось числом"


def test_the_formula_stays_a_formula(handler: XlsxHandler) -> None:
    source = _xlsx_bytes()
    units = handler.extract(source)

    result = handler.apply(source, {units[0].unit_id: "Название проекта"})

    assert "=SUM(B2:B3)" in _cells(result)


def test_ids_are_stable_between_extract_and_apply(handler: XlsxHandler) -> None:
    source = _xlsx_bytes()

    first = [u.unit_id for u in handler.extract(source)]
    second = [u.unit_id for u in handler.extract(handler.apply(source, {}))]

    assert first == second


def test_a_broken_file_is_a_format_error(handler: XlsxHandler) -> None:
    with pytest.raises(DocumentFormatError):
        handler.extract(b"not a workbook")


def test_a_declared_entity_inside_the_workbook_is_refused(handler: XlsxHandler) -> None:
    import zipfile

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(
            "xl/sharedStrings.xml",
            '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "a">]><sst><si><t>&a;</t></si></sst>',
        )
    with pytest.raises(DocumentFormatError):
        handler.extract(buffer.getvalue())
