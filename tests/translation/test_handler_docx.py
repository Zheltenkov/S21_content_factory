"""Обработчик DOCX: подмена текстовых узлов внутри исходного архива.

Переезд существующей механики на общий контракт. Побайтного совпадения тут не требуем — XML
пересобирается, — но требуем смыслового: без переводов все абзацы и всё оформление на месте.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from content_factory.translation.errors import DocumentFormatError
from content_factory.translation.handlers.docx import DocxHandler


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Intro to Machine Learning", level=1)
    paragraph = document.add_paragraph("Summary: ")
    paragraph.add_run("today").bold = True
    paragraph.add_run(" we will help you.")
    document.add_paragraph("")
    document.add_paragraph("1. Chapter I")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _texts(content: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(content)).paragraphs]


@pytest.fixture
def handler() -> DocxHandler:
    return DocxHandler()


def test_paragraphs_with_text_become_units(handler: DocxHandler) -> None:
    units = handler.extract(_docx_bytes())

    assert [u.text for u in units] == [
        "Intro to Machine Learning",
        "Summary: today we will help you.",
        "1. Chapter I",
    ], "пустой абзац не переводим"


def test_without_translations_every_paragraph_survives(handler: DocxHandler) -> None:
    source = _docx_bytes()

    assert _texts(handler.apply(source, {})) == _texts(source)


def test_a_translated_paragraph_replaces_only_itself(handler: DocxHandler) -> None:
    source = _docx_bytes()
    units = handler.extract(source)

    result = handler.apply(source, {units[0].unit_id: "Введение в машинное обучение"})

    texts = _texts(result)
    assert texts[0] == "Введение в машинное обучение"
    assert texts[1] == "Summary: today we will help you.", "соседний абзац не тронут"


def test_inline_styles_survive_the_translation(handler: DocxHandler) -> None:
    source = _docx_bytes()
    units = handler.extract(source)

    result = handler.apply(source, {units[1].unit_id: "Сегодня мы тебе поможем разобраться"})

    runs = Document(BytesIO(result)).paragraphs[1].runs
    assert any(run.bold for run in runs), "жирный run остался жирным"
    assert "".join(run.text for run in runs) == "Сегодня мы тебе поможем разобраться"


def test_ids_are_the_same_on_extract_and_apply(handler: DocxHandler) -> None:
    source = _docx_bytes()

    first = [u.unit_id for u in handler.extract(source)]
    second = [u.unit_id for u in handler.extract(handler.apply(source, {}))]

    assert first == second, "ключи должны быть устойчивы, иначе перевод ляжет не туда"


def test_a_broken_file_is_a_format_error(handler: DocxHandler) -> None:
    with pytest.raises(DocumentFormatError):
        handler.extract(b"not a zip at all")


def test_a_document_that_unpacks_too_large_is_refused(
    handler: DocxHandler, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Загрузка считает сжатый размер, поэтому предел на разворачивание нужен и здесь."""
    import zipfile

    monkeypatch.setattr("content_factory.translation.zip_safety.MAX_MEMBER_BYTES", 2048)
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * 500_000)

    with pytest.raises(DocumentFormatError, match="разворачивается"):
        handler.extract(buffer.getvalue())
