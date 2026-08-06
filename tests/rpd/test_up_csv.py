"""Парсер УП из CSV-экспорта Реестра (путь D — файл-УП)."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from content_factory.rpd.source_fidelity import canonicalize_up_plan
from content_factory.rpd.up_models import from_registry_plan, parse_up_csv

_HEADER = [
    "Тематический блок", "Цели блока", "№", "Название контентной единицы (проекта)", "Краткое описание",
    "Обр. результаты — что узнает (ЗНАТЬ)", "Обр. результаты — что умеет (УМЕТЬ)", "Обр. результаты — какой навык (НАВЫКИ)",
    "Необходимое ПО", "Формат", "Трудоёмкость, астр. часы", "Список навыков (развесовка)",
]


def _write_csv(path: Path) -> None:
    rows = [
        _HEADER,
        ["Модуль А", "Цель А", "1", "Проект 1", "описание1", "знает X", "умеет Y", "навык Z", "Python", "индивидуальный", "10", "Навык Альфа: 30%"],
        ["", "Цель А", "2", "Проект 2", "описание2", "знает W", "", "", "Python", "индивидуальный", "20", "Навык Альфа: 20%"],
        ["Модуль Б", "Цель Б", "3", "Проект 3", "описание3", "", "умеет V", "", "R", "групповой", "5", "Навык Бета"],
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(rows)


def test_parse_csv_groups_projects_into_modules(tmp_path: Path) -> None:
    path = tmp_path / "up.csv"
    _write_csv(path)
    up = parse_up_csv(path, title="Тест")
    assert up.title == "Тест"
    assert [m.title for m in up.modules] == ["Модуль А", "Модуль Б"]
    assert len(up.modules[0].rows) == 2  # блок А удержан для второй строки (пустой блок)
    assert up.modules[0].goals == "Цель А"
    assert up.total_hours == "35"


def test_parse_csv_outcomes_and_skills(tmp_path: Path) -> None:
    path = tmp_path / "up.csv"
    _write_csv(path)
    up = parse_up_csv(path)
    groups = up.outcomes_by_group()
    assert "знает X" in groups["know"] and "умеет Y" in groups["can"]
    codes = [c.code for c in up.competencies()]
    assert "Навык Альфа" in codes and "Навык Бета" in codes  # развесовка «: 30%» отброшена
    assert codes.count("Навык Альфа") == 1  # дедуп


def test_render_from_csv_has_full_columns(tmp_path: Path) -> None:
    import docx

    from content_factory.rpd.render_selective import render_rpd_pk

    path = tmp_path / "up.csv"
    _write_csv(path)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(parse_up_csv(path, title="Тест")))
    doc = docx.Document(str(out))

    def headers(table: object) -> list[str]:
        return [c.text.strip() for c in table.rows[0].cells]  # type: ignore[attr-defined]

    plan = next(t for t in doc.tables if "Наименование модуля" in headers(t))
    assert "Практические занятия" in headers(plan) and "Контроль" in headers(plan)
    comp = next(t for t in doc.tables if headers(t)[0] == "Код")
    assert headers(comp) == [
        "Код",
        "Профессиональная компетенция",
        "Исходный уровень",
        "Качественное изменение",
        "После обучения слушатель способен",
        "Основание",
        "Оценка",
    ]
    modules = next(t for t in doc.tables if "Тема и краткое содержание" in headers(t))
    cells = "\n".join(c.text for r in modules.rows for c in r.cells)
    assert "Проект 1" in cells  # разложено по проектам (имя из УП)


def test_registry_adapter_carries_forward_blank_block_title() -> None:
    up = from_registry_plan(
        {
            "title": "Тест",
            "meta": {"total_hours": 35},
            "rows": [
                {"block_title": "Модуль А", "project_name": "Проект 1", "effort_hours": 10},
                {"block_title": "", "project_name": "Проект 2", "effort_hours": 20},
                {"block_title": "Модуль Б", "project_name": "Проект 3", "effort_hours": 5},
            ],
        }
    )

    assert [module.title for module in up.modules] == ["Модуль А", "Модуль Б"]
    assert [module.hours for module in up.modules] == ["30", "5"]
    assert [len(module.rows) for module in up.modules] == [2, 1]


def test_curriculum_plan_14_golden_source_fidelity() -> None:
    fixture = Path(__file__).parent / "fixtures" / "curriculum_plan_14_golden.csv"
    spec = canonicalize_up_plan(parse_up_csv(fixture, title="curriculum_plan_14"))

    assert spec.module_count == 9
    assert spec.project_count == 24
    assert spec.competency_count == 23
    assert spec.declared_total_hours == "478"
    assert spec.calculated_total_hours == "478"
    assert [module.hours for module in spec.modules] == ["62", "64", "78", "38", "84", "40", "42", "44", "26"]
    assert len(spec.fingerprint) == 64


_REAL = Path(__file__).resolve().parents[2] / "УП_ИИ_предприниматель.csv"


@pytest.mark.skipif(not _REAL.exists(), reason="реальный УП-CSV отсутствует в корне репо")
def test_real_up_csv_parses() -> None:
    up = parse_up_csv(_REAL, title="ИИ для предпринимателя")
    assert len(up.modules) >= 5
    assert up.total_hours == "72"
    assert len(up.competencies()) > 0
    groups = up.outcomes_by_group()
    assert groups["know"] and groups["can"]
