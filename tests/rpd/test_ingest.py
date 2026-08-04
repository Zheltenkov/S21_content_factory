"""ДПО-программа: разбор DOCX в типизированную DpoProgram (Slice 0).

Основные тесты — на синтетическом ДПО-документе (детерминированно, самодостаточно). Плюс один
характеризующий тест на реальном образце в корне репо, который пропускается, если файла нет.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from content_factory.rpd.ingest import parse_dpo_program
from content_factory.rpd.models import Competency, DpoProgram

# ``sample_program`` fixture lives in conftest.py (shared with the transfer tests).


def test_parse_returns_dpo_program_with_sections(sample_program: DpoProgram) -> None:
    assert isinstance(sample_program, DpoProgram)
    numbers = [s.number for s in sample_program.sections]
    assert "1.1" in numbers and "1.12" in numbers


def test_section_lookup_by_number_and_missing_is_none(sample_program: DpoProgram) -> None:
    goal_section = sample_program.section("1.1")
    assert goal_section is not None and "Цель" in goal_section.title
    # раздел отсутствует в документе -> None (двигает gap-вопрос, не выдумку)
    assert sample_program.section("2.3") is None


def test_goal_and_volume_are_transferred_verbatim(sample_program: DpoProgram) -> None:
    assert sample_program.goal.startswith("Программа профессиональной переподготовки")
    assert "300" in sample_program.volume_hours


def test_planned_results_are_extracted_as_lines(sample_program: DpoProgram) -> None:
    results = sample_program.planned_results
    assert any("приобретает знания" in line for line in results)
    assert any("осваивает методы" in line for line in results)


def test_competencies_come_from_the_1_12_table(sample_program: DpoProgram) -> None:
    comps = sample_program.competencies
    assert all(isinstance(c, Competency) for c in comps)
    codes = [c.code for c in comps]
    assert "ПК - 1" in codes and len(comps) == 2
    pk1 = next(c for c in comps if c.code == "ПК - 1")
    assert "аналитическую" in pk1.description
    # заголовок таблицы не попадает в компетенции
    assert "Компетенции" not in codes


def test_absent_slot_yields_empty_not_invention(tmp_path: Path) -> None:
    # программа без раздела «Объем» -> volume_hours пустой, а не выдуманный
    path = tmp_path / "bare.docx"
    d = docx.Document()
    d.add_heading("1.1. Цель реализации программы", level=2)
    d.add_paragraph("Некоторая цель.")
    d.save(str(path))
    program = parse_dpo_program(path)
    assert program.volume_hours == ""
    assert program.competencies == []
    assert program.planned_results == []


_REAL = Path(__file__).resolve().parents[2] / "Биоинформатика_и_искусственный_интеллект_в_медицине_итог.docx"


@pytest.mark.skipif(not _REAL.exists(), reason="реальный образец ДПО отсутствует в корне репо")
def test_real_sample_is_parsed_end_to_end() -> None:
    program = parse_dpo_program(_REAL)
    assert program.goal.startswith("Программа")
    assert "300" in program.volume_hours
    codes = [c.code for c in program.competencies]
    assert any(code.startswith("ПК") for code in codes)
    assert len(program.planned_results) >= 3
    # ключевые разделы характеристики распознаны
    assert program.section("1.1") is not None
    assert program.section("1.11") is not None
