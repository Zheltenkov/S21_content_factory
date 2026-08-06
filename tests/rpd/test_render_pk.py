"""Выборочная сборка РПД на реальном ДПО-шаблоне ПК (B1).

Проверяем: УП-разделы заполнены; boilerplate (нормативка) сохранён; проза/скрины/пробелы помечены.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import docx

from content_factory.rpd.dpo_spec import Cat, category_for
from content_factory.rpd.render_selective import render_rpd_pk
from content_factory.rpd.template import default_pk_skeleton_path
from content_factory.rpd.transfer import GAP_MARKER
from content_factory.rpd.up_models import from_registry_plan

_PLAN = {
    "title": "Тестовый УП",
    "meta": {"total_hours": 300},
    "rows": [
        {"block_title": "Модуль 1", "outcomes_know": "знает основы", "outcomes_can": "умеет применять", "outcomes_skills": "", "effort_hours": 150, "skills": ["Анализ данных"]},
        {"block_title": "Модуль 2", "outcomes_know": "", "outcomes_can": "строит модели", "outcomes_skills": "", "effort_hours": 150, "skills": ["Моделирование"]},
    ],
}


def _render(tmp_path: Path) -> tuple[str, str]:
    up = from_registry_plan(_PLAN)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(up))
    doc = docx.Document(str(out))
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    return paragraphs, cells


def test_category_for_key_sections() -> None:
    assert category_for("3.1", "Учебный план") is Cat.UP_PLAN
    assert category_for("2.1", "Профессиональные компетенции") is Cat.UP_COMP
    assert category_for("1.3", "Актуальность программы") is Cat.GENERATE
    assert category_for("1.2", "Нормативно-правовые основания") is Cat.BOILERPLATE
    assert category_for("3.3.1", "Пример оформления учебного модуля") is Cat.UP_MODULES
    assert category_for("4.3", "Паспорт оценочных материалов") is Cat.GAP


def test_render_is_valid_docx(tmp_path: Path) -> None:
    up = from_registry_plan(_PLAN)
    data = render_rpd_pk(up)
    assert docx.Document(__import__("io").BytesIO(data)) is not None


def test_up_sections_are_filled(tmp_path: Path) -> None:
    _paras, cells = _render(tmp_path)
    assert "Модуль 1" in cells and "Модуль 2" in cells  # учебный план из УП
    assert "150" in cells  # часы модуля
    assert "Анализ данных" in cells  # компетенции из skills


def test_boilerplate_is_kept(tmp_path: Path) -> None:
    paragraphs, _cells = _render(tmp_path)
    assert "273-ФЗ" in paragraphs  # нормативка ФОРМЫ сохранена, не затёрта


def test_prose_and_platform_and_gaps_are_marked(tmp_path: Path) -> None:
    paragraphs, _cells = _render(tmp_path)
    assert "[Будет сгенерировано]" in paragraphs  # актуальность/цели — под генерацию (без LLM)
    assert "скрин" in paragraphs.lower()  # примеры/оценка — скрин с платформы
    assert "Методист" in paragraphs  # оценочные/разработчики — пометка методисту


def test_no_comments_leak_into_output(tmp_path: Path) -> None:
    import io
    import zipfile

    data = render_rpd_pk(from_registry_plan(_PLAN))
    archive = zipfile.ZipFile(io.BytesIO(data))
    assert "word/comments.xml" not in archive.namelist()  # заметки методиста вырезаны
    xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    assert "commentReference" not in xml and "commentRangeStart" not in xml


def test_render_preserves_template_sections_and_orientation() -> None:
    template = docx.Document(str(default_pk_skeleton_path()))
    rendered = docx.Document(BytesIO(render_rpd_pk(from_registry_plan(_PLAN))))

    assert len(rendered.sections) == len(template.sections) == 4
    assert [section.orientation for section in rendered.sections] == [
        section.orientation for section in template.sections
    ]


def test_module_program_combines_up_rows_with_platform_evidence_gate() -> None:
    rendered = docx.Document(BytesIO(render_rpd_pk(from_registry_plan(_PLAN))))
    module_table = next(
        table
        for table in rendered.tables
        if table.rows and "Тема и краткое содержание" in [cell.text.strip() for cell in table.rows[0].cells]
    )
    text = "\n".join(cell.text for row in module_table.rows for cell in row.cells)

    assert "Модуль 1" in text and "Модуль 2" in text
    assert text.count("вставьте скрин заданий и контролей модуля") == 2


def test_empty_up_keeps_boilerplate_and_gaps_up_sections(tmp_path: Path) -> None:
    from content_factory.rpd.up_models import UpPlan

    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(UpPlan()))
    doc = docx.Document(str(out))
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    assert "273-ФЗ" in paragraphs  # boilerplate всё равно на месте
    assert GAP_MARKER in paragraphs  # УП-разделы без данных → пробел
