"""Template-profile contract, recognition and current ПК golden semantics."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from content_factory.rpd.build import build_rpd_sections
from content_factory.rpd.template_profiles import (
    BUILTIN_PK_PROFILE_ID,
    UnsupportedTemplateProfileError,
    detect_template_profile,
    get_template_profile,
    registered_template_profiles,
)
from content_factory.rpd.up_models import from_registry_plan


def _plan() -> dict:
    return {
        "title": "Профильный тест",
        "meta": {"total_hours": 16},
        "rows": [
            {
                "block_title": "Модуль",
                "project_name": "Проект",
                "effort_hours": 16,
                "skills": ["Навык"],
            }
        ],
    }


def test_builtin_pk_profile_satisfies_its_capability_contract() -> None:
    profile = get_template_profile(BUILTIN_PK_PROFILE_ID)

    report = profile.inspect(profile.template_path())

    assert report.supported
    assert report.program_type == "повышение квалификации"
    assert report.section_count > 10
    assert len(report.template_sha256) == 64
    assert not report.missing_bindings
    assert not report.missing_table_labels
    detected, detected_report = detect_template_profile(profile.template_path())
    assert detected.profile_id == BUILTIN_PK_PROFILE_ID
    assert detected_report.template_sha256 == report.template_sha256


def test_build_persists_profile_and_template_fingerprints() -> None:
    build = build_rpd_sections(from_registry_plan(_plan()))

    assert build.template_profile_id == BUILTIN_PK_PROFILE_ID
    assert build.template_report is not None and build.template_report.supported
    assert build.generation_trace.template_sha256 == build.template_report.template_sha256
    assert len(build.generation_trace.source_sha256) == 64


def test_profile_rejects_docx_without_required_bindings(tmp_path: Path) -> None:
    path = tmp_path / "unknown.docx"
    document = docx.Document()
    document.add_heading("Произвольный документ", level=1)
    document.add_paragraph("Нет контракта РПД.")
    document.save(path)

    report = get_template_profile().inspect(path)

    assert not report.supported
    assert "program_title" in report.missing_bindings
    assert "Категория слушателей" in report.missing_table_labels
    with pytest.raises(UnsupportedTemplateProfileError, match="Unsupported RPD template"):
        detect_template_profile(path)


def test_unknown_profile_fails_explicitly() -> None:
    assert registered_template_profiles() == (BUILTIN_PK_PROFILE_ID,)
    with pytest.raises(UnsupportedTemplateProfileError, match="university_v1"):
        get_template_profile("university_v1")
