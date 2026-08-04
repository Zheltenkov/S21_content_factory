"""Shared ДПО-fixtures for the rpd tests."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from content_factory.rpd.ingest import parse_dpo_program
from content_factory.rpd.models import DpoProgram
from content_factory.rpd.template import RpdTemplate, load_rpd_template


def build_sample_dpo(path: Path) -> None:
    """Собирает минимальную ДПО-программу в формате разделов Школы 21.

    Намеренно НЕ содержит разделы 1.3 (актуальность) и 2.3 (рабочая программа) — чтобы тесты
    проверяли толерантность к пробелам и корректную генерацию gap-вопросов вместо выдумки.
    """
    d = docx.Document()
    d.add_heading("1. ХАРАКТЕРИСТИКА ПРОГРАММЫ", level=1)
    d.add_heading("1.1. Цель реализации программы", level=2)
    d.add_paragraph("Программа профессиональной переподготовки направлена на получение компетенций.")
    d.add_heading("1.6. Объем программы:", level=2)
    d.add_paragraph("300 академических часов.")
    d.add_heading("1.11. Планируемые результаты обучения", level=2)
    d.add_paragraph("В результате освоения программы обучающийся:")
    d.add_paragraph("приобретает знания в области анализа данных;")
    d.add_paragraph("осваивает методы обработки данных;")
    d.add_heading("1.12. Профессиональные компетенции", level=2)
    d.add_paragraph("Выпускник обладает следующими компетенциями:")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Компетенции"
    table.rows[0].cells[1].text = "Описание компетенции"
    for code, desc in [("ПК - 1", "Способен осуществлять аналитическую деятельность"), ("ПК - 2", "Способен применять методы анализа")]:
        row = table.add_row()
        row.cells[0].text = code
        row.cells[1].text = desc
    d.add_heading("2. СТРУКТУРА И СОДЕРЖАНИЕ ПРОГРАММЫ", level=1)
    d.save(str(path))


@pytest.fixture
def sample_program(tmp_path: Path) -> DpoProgram:
    path = tmp_path / "dpo.docx"
    build_sample_dpo(path)
    return parse_dpo_program(path)


def build_skeleton(path: Path) -> None:
    """Каркас-шаблон РПД: как sample, но с ДОПОЛНИТЕЛЬНЫМ разделом 1.3, которого нет в программе.

    Все тела помечены «ПРИМЕР» — тесты проверяют, что этот текст образца не протекает в выход.
    """
    d = docx.Document()
    d.add_heading("1. ХАРАКТЕРИСТИКА ПРОГРАММЫ", level=1)
    d.add_heading("1.1. Цель реализации программы", level=2)
    d.add_paragraph("ПРИМЕР: цель из образца — не должна утечь в выход.")
    d.add_heading("1.3. Актуальность программы", level=2)
    d.add_paragraph("ПРИМЕР: актуальность из образца.")
    d.add_heading("1.6. Объем программы:", level=2)
    d.add_paragraph("ПРИМЕР: 999 часов.")
    d.add_heading("1.11. Планируемые результаты обучения", level=2)
    d.add_paragraph("ПРИМЕР: результаты.")
    d.add_heading("1.12. Профессиональные компетенции", level=2)
    d.add_paragraph("ПРИМЕР.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Компетенции"
    table.rows[0].cells[1].text = "Описание компетенции"
    row = table.add_row()
    row.cells[0].text = "ПК - 9"
    row.cells[1].text = "ПРИМЕР"
    d.add_heading("2. СТРУКТУРА И СОДЕРЖАНИЕ ПРОГРАММЫ", level=1)
    d.add_paragraph("ПРИМЕР: структура.")
    d.save(str(path))


@pytest.fixture
def skeleton_path(tmp_path: Path) -> Path:
    path = tmp_path / "skeleton.docx"
    build_skeleton(path)
    return path


@pytest.fixture
def template(skeleton_path: Path) -> RpdTemplate:
    return load_rpd_template(skeleton_path)
