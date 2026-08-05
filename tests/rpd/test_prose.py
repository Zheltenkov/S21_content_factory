"""LLM-генерация прозы (актуальность/цели/паспорт) и её вставка в рендер (B3)."""

from __future__ import annotations

from pathlib import Path

import docx

from content_factory.rpd.prose import generate_prose
from content_factory.rpd.render_selective import render_rpd_pk
from content_factory.rpd.up_models import UpPlan, from_registry_plan

_PLAN = {
    "title": "Тестовая программа",
    "meta": {"total_hours": 72},
    "rows": [{"block_title": "Модуль 1", "outcomes_know": "знает основы", "effort_hours": 72, "skills": ["Навык"]}],
}


class _FakeClient:
    def __init__(self, text: str) -> None:
        self.text = text
        self.calls: list[tuple[str, str]] = []

    def complete(self, system: str, user: str, **kwargs: object) -> str:
        del kwargs
        self.calls.append((system, user))
        return self.text


def test_generate_prose_returns_paragraphs_for_prose_section() -> None:
    up = from_registry_plan(_PLAN)
    client = _FakeClient("Актуальность высока.\nВторой абзац актуальности.")
    paragraphs = generate_prose("Актуальность программы", up, client)
    assert paragraphs == ["Актуальность высока.", "Второй абзац актуальности."]
    assert client.calls and "Тестовая программа" in client.calls[0][1]  # УП-контекст в промпте


def test_generate_prose_skips_non_prose_and_missing_client() -> None:
    up = from_registry_plan(_PLAN)
    assert generate_prose("Учебный план", up, _FakeClient("x")) == []  # не проза
    assert generate_prose("Актуальность программы", up, None) == []  # нет клиента


def test_render_uses_generated_prose_instead_of_placeholder(tmp_path: Path) -> None:
    up = from_registry_plan(_PLAN)
    client = _FakeClient("Сгенерированная проза.")
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(up, llm_client=client))
    paragraphs = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Сгенерированная проза." in paragraphs
    assert "[Будет сгенерировано]" not in paragraphs  # все GENERATE-разделы заполнены


def test_render_without_client_falls_back_to_placeholder(tmp_path: Path) -> None:
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(from_registry_plan(_PLAN)))
    paragraphs = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[Будет сгенерировано]" in paragraphs  # без клиента — плейсхолдер


def test_render_with_failing_client_falls_back(tmp_path: Path) -> None:
    class _Boom:
        def complete(self, system: str, user: str, **kwargs: object) -> str:
            raise RuntimeError("llm down")

    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_pk(UpPlan(title="X", modules=[]), llm_client=_Boom()))
    # даже при сбое LLM документ собирается, проза → плейсхолдер
    paragraphs = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[Будет сгенерировано]" in paragraphs
