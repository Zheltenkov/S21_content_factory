"""Сборка РПД в DOCX из черновика переноса (Slice 2).

Каркас даёт стили и структуру; тело пересобирается контентом программы. Ключевая гарантия — no-leak:
текст образца-каркаса («ПРИМЕР…») в выход не попадает; отсутствующие разделы помечены маркером.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from content_factory.rpd.ingest import parse_dpo_program
from content_factory.rpd.models import DpoProgram
from content_factory.rpd.render_docx import render_rpd_docx
from content_factory.rpd.template import RpdTemplate, load_rpd_template
from content_factory.rpd.transfer import GAP_MARKER, build_rpd


def _all_text(path: Path) -> str:
    parsed = parse_dpo_program(path)
    return "\n".join(section.text for section in parsed.sections)


def test_render_writes_a_valid_docx(sample_program: DpoProgram, template: RpdTemplate, skeleton_path: Path, tmp_path: Path) -> None:
    draft = build_rpd(sample_program, template)
    data = render_rpd_docx(skeleton_path, draft, program_title="Тестовая программа")
    out = tmp_path / "rpd.docx"
    out.write_bytes(data)
    # открывается как валидный docx
    doc = docx.Document(str(out))
    assert len(doc.paragraphs) > 0


def test_render_carries_program_content_not_example(sample_program: DpoProgram, template: RpdTemplate, skeleton_path: Path, tmp_path: Path) -> None:
    draft = build_rpd(sample_program, template)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_docx(skeleton_path, draft, program_title="Тестовая программа"))
    text = _all_text(out)
    assert "Программа профессиональной переподготовки" in text  # перенос цели из программы
    assert "300" in text  # объём из программы
    assert "ПРИМЕР" not in text  # текст каркаса-образца не протёк (no-leak)


def test_render_marks_missing_sections(sample_program: DpoProgram, template: RpdTemplate, skeleton_path: Path, tmp_path: Path) -> None:
    draft = build_rpd(sample_program, template)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_docx(skeleton_path, draft, program_title="Тестовая программа"))
    assert GAP_MARKER in _all_text(out)  # 1.3 / 2 помечены пробелом


def test_render_carries_competency_table(sample_program: DpoProgram, template: RpdTemplate, skeleton_path: Path, tmp_path: Path) -> None:
    draft = build_rpd(sample_program, template)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_docx(skeleton_path, draft, program_title="Тестовая программа"))
    doc = docx.Document(str(out))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert any("ПК - 1" in c for c in cells)  # компетенция перенесена
    assert not any("ПК - 9" in c for c in cells)  # пример из каркаса не протёк


def test_render_empty_program_is_all_gaps_no_leak(template: RpdTemplate, skeleton_path: Path, tmp_path: Path) -> None:
    draft = build_rpd(DpoProgram(), template)
    out = tmp_path / "rpd.docx"
    out.write_bytes(render_rpd_docx(skeleton_path, draft))
    text = _all_text(out)
    assert GAP_MARKER in text
    assert "ПРИМЕР" not in text  # даже при пустой программе образец не протекает


_REAL = Path(__file__).resolve().parents[2] / "Биоинформатика_и_искусственный_интеллект_в_медицине_итог.docx"


@pytest.mark.skipif(not _REAL.exists(), reason="реальный образец ДПО отсутствует в корне репо")
def test_render_real_sample_round_trips(tmp_path: Path) -> None:
    template = load_rpd_template(_REAL)
    program = parse_dpo_program(_REAL)
    draft = build_rpd(program, template)
    out = tmp_path / "rpd_real.docx"
    out.write_bytes(render_rpd_docx(_REAL, draft, program_title="Биоинформатика и ИИ в медицине"))
    text = _all_text(out)
    assert "биоинформатик" in text.lower()  # реальный контент дошёл
    assert len(text) > 500
