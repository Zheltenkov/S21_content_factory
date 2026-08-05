"""РПД-режим (web): загрузка программы → сборка РПД → скачивание + персист. Реальный E2E-движок."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from content_factory.api.routers import rpd


def _program_docx() -> bytes:
    d = docx.Document()
    d.add_heading("1.1. Цель реализации программы", level=2)
    d.add_paragraph("Программа профессиональной переподготовки по тестовой теме.")
    d.add_heading("1.12. Профессиональные компетенции", level=2)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Компетенции"
    table.rows[0].cells[1].text = "Описание компетенции"
    row = table.add_row()
    row.cells[0].text = "ПК - 1"
    row.cells[1].text = "Способен тестировать"
    buffer = io.BytesIO()
    d.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> TestClient:
    monkeypatch.setenv("DISABLE_AUTH", "true")
    monkeypatch.setattr(rpd, "RPD_ROOT", tmp_path / "rpd")
    rpd._UPLOADS.clear()
    rpd._RUNS.clear()
    rpd._LOADED = False
    app = FastAPI()
    app.include_router(rpd.page_router)
    return TestClient(app)


def _upload(client: TestClient) -> str:
    resp = client.post("/app/rpd/upload", files={"program": ("prog.docx", _program_docx(),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert resp.status_code == 200
    return resp.json()["upload_id"]


def _generate(client: TestClient) -> str:
    upload_id = _upload(client)
    resp = client.post("/app/rpd/generate", data={"upload_id": upload_id})
    assert resp.status_code == 200
    return resp.json()["run_id"]


def test_home_renders(client: TestClient) -> None:
    resp = client.get("/app/rpd")
    assert resp.status_code == 200
    assert "РПД из программы" in resp.text


def test_upload_requires_a_docx(client: TestClient) -> None:
    assert client.post("/app/rpd/upload").status_code == 400
    bad = client.post("/app/rpd/upload", files={"program": ("x.txt", b"hi", "text/plain")})
    assert bad.status_code == 400


def test_generate_builds_a_run_and_view_shows_coverage(client: TestClient) -> None:
    run_id = _generate(client)
    view = client.get(f"/app/rpd/{run_id}")
    assert view.status_code == 200
    assert "готовность переноса" in view.text
    assert "Скачать РПД" in view.text


def test_generate_with_bad_upload_id_is_rejected(client: TestClient) -> None:
    assert client.post("/app/rpd/generate", data={"upload_id": "nope"}).status_code == 400


def test_download_returns_a_valid_docx(client: TestClient) -> None:
    run_id = _generate(client)
    resp = client.get(f"/app/rpd/{run_id}/download")
    assert resp.status_code == 200
    assert zipfile.is_zipfile(io.BytesIO(resp.content))  # docx = zip
    # содержимое программы дошло, пример каркаса — нет
    with zipfile.ZipFile(io.BytesIO(resp.content)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    assert "переподготовки по тестовой теме" in xml
    assert "Биоинформатик" not in xml  # каркас-образец не протёк


def test_delete_removes_the_run(client: TestClient) -> None:
    run_id = _generate(client)
    assert client.get(f"/app/rpd/{run_id}").status_code == 200
    deleted = client.post(f"/app/rpd/{run_id}/delete", follow_redirects=False)
    assert deleted.status_code == 303
    assert client.get(f"/app/rpd/{run_id}").status_code == 404


def test_runs_survive_reload_and_list_in_recent(client: TestClient) -> None:
    run_id = _generate(client)
    # имитируем рестарт: сбрасываем in-memory индекс и флаг загрузки
    rpd._RUNS.clear()
    rpd._LOADED = False
    home = client.get("/app/rpd")
    assert home.status_code == 200
    assert "Недавние РПД" in home.text
    # переоткрытие после reload работает
    assert client.get(f"/app/rpd/{run_id}").status_code == 200
