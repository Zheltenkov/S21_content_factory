"""Редактируемый черновик РПД (HITL H1): build (УП→разделы) и render_build (черновик→DOCX)."""

from __future__ import annotations

import io
import zipfile

import docx

from content_factory.rpd.build import RpdBuild, build_rpd_sections
from content_factory.rpd.render_selective import render_build
from content_factory.rpd.up_models import from_registry_plan

_PLAN = {
    "title": "Тестовая программа",
    "meta": {"total_hours": 72},
    "rows": [
        {"block_title": "Модуль 1", "outcomes_know": "знает основы", "outcomes_can": "умеет применять", "effort_hours": 36, "skills": ["Навык А"]},
        {"block_title": "Модуль 2", "outcomes_can": "строит модели", "effort_hours": 36, "skills": ["Навык Б"]},
    ],
}


def test_build_produces_editable_sections() -> None:
    build = build_rpd_sections(from_registry_plan(_PLAN))
    assert isinstance(build, RpdBuild)
    plan = build.section("3.1", "Учебный план")
    assert plan is not None and plan.kind == "table" and plan.label == "из УП"
    assert plan.table_header[:2] == ["№", "Наименование модуля"]
    assert any("Модуль 1" in row for row in plan.table_rows)
    # проза без клиента → плейсхолдер, но раздел редактируемый
    relevance = build.section("1.3", "Актуальность программы")
    assert relevance is not None and relevance.kind == "paragraphs" and relevance.label == "сгенерировано"
    # boilerplate в черновик не попадает (его держит шаблон)
    assert build.section("1.2", "Нормативно-правовые основания") is None


def test_render_build_reflects_edits() -> None:
    build = build_rpd_sections(from_registry_plan(_PLAN))
    relevance = build.section("1.3", "Актуальность программы")
    assert relevance is not None
    relevance.paragraphs = ["Отредактированный методистом текст актуальности."]
    data = render_build(build)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8", errors="ignore")
    assert "Отредактированный методистом текст" in xml  # правка попала в DOCX
    assert "273-ФЗ" in xml  # boilerplate ФОРМЫ на месте
    assert "word/comments.xml" not in zipfile.ZipFile(io.BytesIO(data)).namelist()


def test_render_build_table_edit_flows_through() -> None:
    build = build_rpd_sections(from_registry_plan(_PLAN))
    plan = build.section("3.1", "Учебный план")
    assert plan is not None
    plan.table_rows[0][1] = "ПЕРЕИМЕНОВАННЫЙ МОДУЛЬ"
    data = render_build(build)
    doc = docx.Document(io.BytesIO(data))
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "ПЕРЕИМЕНОВАННЫЙ МОДУЛЬ" in cells
